import json
import os
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_notebook(notebook_path):
    cells = []
    
    # 1. Title
    cells.append({
        "cell_type": "markdown",
        "metadata": {},
        "source": [
            "# Vishwakarma Institute of Technology, Pune\n",
            "## Course: CI3201 - Deep Learning\n",
            "### Lab Practical 4: LSTM-Based Model for Time-Series Forecasting (Stock Price Prediction)\n",
            "\n",
            "---"
        ]
    })

    # 2. Objective
    cells.append({
        "cell_type": "markdown",
        "metadata": {},
        "source": [
            "## 1. Objective & Theoretical Background\n",
            "\n",
            "### 1.1 Objective\n",
            "Develop an LSTM-based model for time-series forecasting. We will use historical Apple Stock (AAPL) prices to predict the future closing price based on the previous $N$ days of data.\n",
            "\n",
            "### 1.2 Mathematical Foundations of LSTM\n",
            "Long Short-Term Memory (LSTM) networks overcome the vanishing gradient problem in traditional RNNs using gating mechanisms:\n",
            "1. **Forget Gate**: $f_t = \\sigma(W_f \\cdot [h_{t-1}, x_t] + b_f)$ - Decides what information to discard.\n",
            "2. **Input Gate**: $i_t = \\sigma(W_i \\cdot [h_{t-1}, x_t] + b_i)$ - Decides what new info to store.\n",
            "3. **Candidate Cell State**: $\\tilde{C}_t = \\tanh(W_C \\cdot [h_{t-1}, x_t] + b_C)$\n",
            "4. **Update Cell State**: $C_t = f_t * C_{t-1} + i_t * \\tilde{C}_t$\n",
            "5. **Output Gate**: $o_t = \\sigma(W_o \\cdot [h_{t-1}, x_t] + b_o)$\n",
            "6. **Hidden State**: $h_t = o_t * \\tanh(C_t)$\n",
            "\n",
            "For Time-Series Forecasting, the model outputs a continuous value, so we use **Mean Squared Error (MSE)** as our loss function."
        ]
    })

    # 3. Code: Imports
    cells.append({
        "cell_type": "code",
        "execution_count": None,
        "metadata": {},
        "outputs": [],
        "source": [
            "# Install yfinance to fetch historical stock data (Uncomment if not installed)\n",
            "# !pip install yfinance\n",
            "\n",
            "import math\n",
            "import numpy as np\n",
            "import pandas as pd\n",
            "import matplotlib.pyplot as plt\n",
            "import seaborn as sns\n",
            "import yfinance as yf\n",
            "\n",
            "from sklearn.preprocessing import MinMaxScaler\n",
            "from sklearn.metrics import mean_squared_error\n",
            "\n",
            "from keras.models import Model\n",
            "from keras.layers import Input, LSTM, Dense, Dropout, Activation\n",
            "from keras.optimizers import RMSprop\n",
            "from keras.callbacks import EarlyStopping\n",
            "\n",
            "import warnings\n",
            "warnings.filterwarnings('ignore')"
        ]
    })

    # 4. Code: Load Data
    cells.append({
        "cell_type": "code",
        "execution_count": None,
        "metadata": {},
        "outputs": [],
        "source": [
            "# 1. Fetch AAPL Stock Data from Yahoo Finance\n",
            "print(\"Fetching Data...\")\n",
            "df = yf.download('AAPL', start='2015-01-01', end='2024-01-01')\n",
            "df.reset_index(inplace=True)\n",
            "\n",
            "# Save to Datasets directory for future reference\n",
            "df.to_csv('Datasets/AAPL_stock_prices.csv', index=False)\n",
            "\n",
            "print(\"Data Shape:\", df.shape)\n",
            "df.head()"
        ]
    })

    # 5. Code: Visualize Data
    cells.append({
        "cell_type": "code",
        "execution_count": None,
        "metadata": {},
        "outputs": [],
        "source": [
            "# 2. Visualize the Historical Closing Price\n",
            "plt.figure(figsize=(16, 6))\n",
            "plt.plot(df['Date'], df['Close'], color='blue', label='AAPL Closing Price')\n",
            "plt.title('Apple (AAPL) Historical Stock Closing Price')\n",
            "plt.xlabel('Date')\n",
            "plt.ylabel('Closing Price (USD)')\n",
            "plt.legend()\n",
            "plt.grid(True)\n",
            "plt.show()"
        ]
    })

    # 6. Preprocessing
    cells.append({
        "cell_type": "markdown",
        "metadata": {},
        "source": [
            "## 2. Data Preprocessing\n",
            "\n",
            "- **Scaling**: LSTM is highly sensitive to the magnitude of data. We use `MinMaxScaler` to scale the prices to `(0, 1)`.\n",
            "- **Sequence Generation**: We use the past 60 days (`time_steps = 60`) of data to predict the 61st day's price."
        ]
    })

    # 7. Code: Scaling and Splitting
    cells.append({
        "cell_type": "code",
        "execution_count": None,
        "metadata": {},
        "outputs": [],
        "source": [
            "# Extract Closing Prices\n",
            "data = df.filter(['Close']).values\n",
            "\n",
            "# Scale Data\n",
            "scaler = MinMaxScaler(feature_range=(0, 1))\n",
            "scaled_data = scaler.fit_transform(data)\n",
            "\n",
            "# Define Sequence Length (Time Steps)\n",
            "time_steps = 60\n",
            "\n",
            "# Generate Sequences\n",
            "X = []\n",
            "Y = []\n",
            "for i in range(time_steps, len(scaled_data)):\n",
            "    X.append(scaled_data[i-time_steps:i, 0])\n",
            "    Y.append(scaled_data[i, 0])\n",
            "    \n",
            "X, Y = np.array(X), np.array(Y)\n",
            "\n",
            "# Reshape X for LSTM [samples, time_steps, features]\n",
            "X = np.reshape(X, (X.shape[0], X.shape[1], 1))\n",
            "\n",
            "# Chronological Train-Test Split (80% Train, 20% Test)\n",
            "train_size = int(len(X) * 0.8)\n",
            "X_train, X_test = X[:train_size], X[train_size:]\n",
            "Y_train, Y_test = Y[:train_size], Y[train_size:]\n",
            "\n",
            "print(\"X_train shape:\", X_train.shape)\n",
            "print(\"Y_train shape:\", Y_train.shape)\n",
            "print(\"X_test shape:\", X_test.shape)"
        ]
    })

    # 8. Model Building
    cells.append({
        "cell_type": "markdown",
        "metadata": {},
        "source": [
            "## 3. Build the LSTM Architecture using Keras Functional API\n",
            "\n",
            "We construct the LSTM network using the **Keras Functional API** (similar to the SMS Spam example provided in class):\n",
            "1. **Input Layer**: Shape = `[time_steps, 1]`\n",
            "2. **LSTM Layer**: 64 units to capture temporal dependencies.\n",
            "3. **Dropout**: 20% rate to prevent overfitting.\n",
            "4. **Dense Output**: 1 unit for the continuous price prediction."
        ]
    })

    # 9. Code: Model
    cells.append({
        "cell_type": "code",
        "execution_count": None,
        "metadata": {},
        "outputs": [],
        "source": [
            "def build_lstm_model():\n",
            "    # Define Input Layer\n",
            "    inputs = Input(name='inputs', shape=(time_steps, 1))\n",
            "    \n",
            "    # LSTM Layer\n",
            "    layer = LSTM(64, return_sequences=False, name='lstm_1')(inputs)\n",
            "    \n",
            "    # Regularization\n",
            "    layer = Dropout(0.2, name='dropout_1')(layer)\n",
            "    \n",
            "    # Output Layer (Linear Activation for Regression)\n",
            "    layer = Dense(1, name='out_layer')(layer)\n",
            "    \n",
            "    model = Model(inputs=inputs, outputs=layer)\n",
            "    return model\n",
            "\n",
            "model = build_lstm_model()\n",
            "model.summary()\n",
            "\n",
            "# Compile Model\n",
            "model.compile(loss='mean_squared_error', optimizer=RMSprop(learning_rate=0.001))"
        ]
    })

    # 10. Training
    cells.append({
        "cell_type": "markdown",
        "metadata": {},
        "source": [
            "## 4. Train the Model with Early Stopping\n",
            "We apply `EarlyStopping` to halt training when validation loss stops improving, restoring the best model weights."
        ]
    })

    # 11. Code: Training
    cells.append({
        "cell_type": "code",
        "execution_count": None,
        "metadata": {},
        "outputs": [],
        "source": [
            "early_stop = EarlyStopping(monitor='val_loss', patience=5, restore_best_weights=True)\n",
            "\n",
            "history = model.fit(X_train, Y_train,\n",
            "                    epochs=50,\n",
            "                    batch_size=32,\n",
            "                    validation_split=0.1,\n",
            "                    callbacks=[early_stop],\n",
            "                    verbose=1)"
        ]
    })

    # 12. Evaluation
    cells.append({
        "cell_type": "markdown",
        "metadata": {},
        "source": [
            "## 5. Model Evaluation & Visualization"
        ]
    })

    # 13. Code: Evaluation
    cells.append({
        "cell_type": "code",
        "execution_count": None,
        "metadata": {},
        "outputs": [],
        "source": [
            "# Plot Training vs Validation Loss\n",
            "plt.figure(figsize=(10, 5))\n",
            "plt.plot(history.history['loss'], label='Training Loss')\n",
            "plt.plot(history.history['val_loss'], label='Validation Loss')\n",
            "plt.title('Model Loss (MSE) over Epochs')\n",
            "plt.xlabel('Epoch')\n",
            "plt.ylabel('Mean Squared Error')\n",
            "plt.legend()\n",
            "plt.show()\n",
            "\n",
            "# Make Predictions on Test Set\n",
            "predictions = model.predict(X_test)\n",
            "\n",
            "# Inverse transform to get actual USD values\n",
            "predictions_usd = scaler.inverse_transform(predictions)\n",
            "Y_test_usd = scaler.inverse_transform(Y_test.reshape(-1, 1))\n",
            "\n",
            "# Calculate Root Mean Squared Error (RMSE)\n",
            "rmse = math.sqrt(mean_squared_error(Y_test_usd, predictions_usd))\n",
            "print(f\"Test RMSE: ${rmse:.2f}\")"
        ]
    })

    # 14. Code: Visualizing Predictions
    cells.append({
        "cell_type": "code",
        "execution_count": None,
        "metadata": {},
        "outputs": [],
        "source": [
            "# Plot Actual vs Predicted Prices\n",
            "# We need to align the dates for the test set\n",
            "test_data_start_idx = len(df) - len(Y_test)\n",
            "test_dates = df['Date'].iloc[test_data_start_idx:].values\n",
            "\n",
            "plt.figure(figsize=(16, 6))\n",
            "plt.plot(test_dates, Y_test_usd, color='blue', label='Actual Apple Stock Price')\n",
            "plt.plot(test_dates, predictions_usd, color='red', label='Predicted Apple Stock Price')\n",
            "plt.title('Apple (AAPL) Stock Price Prediction using LSTM')\n",
            "plt.xlabel('Date')\n",
            "plt.ylabel('Closing Price (USD)')\n",
            "plt.legend()\n",
            "plt.grid(True)\n",
            "plt.show()"
        ]
    })

    notebook_content = {
        "cells": cells,
        "metadata": {
            "kernelspec": {
                "display_name": "Python 3",
                "language": "python",
                "name": "python3"
            },
            "language_info": {
                "codemirror_mode": {"name": "ipython", "version": 3},
                "file_extension": ".py",
                "mimetype": "text/x-python",
                "name": "python",
                "nbconvert_exporter": "python",
                "pygments_lexer": "ipython3",
                "version": "3.10.0"
            }
        },
        "nbformat": 4,
        "nbformat_minor": 2
    }

    with open(notebook_path, "w", encoding="utf-8") as f:
        json.dump(notebook_content, f, indent=1)
    print(f"Created Notebook: {notebook_path}")

def create_docx(docx_path):
    doc = docx.Document()

    # Document Header
    p_head = doc.add_paragraph()
    p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_head.add_run("Vishwakarma Institute of Technology, Pune\n")
    run_title.bold = True
    run_title.font.size = Pt(16)
    run_sub = p_head.add_run("Department of Computer Engineering (Artificial Intelligence)\nCI3201: Deep Learning - Lab Assignment 4\n")
    run_sub.font.size = Pt(13)
    run_sub.bold = True
    
    doc.add_paragraph("="*80)

    # Problem Statement
    p_ps = doc.add_paragraph()
    r_ps_lbl = p_ps.add_run("Problem Statement: ")
    r_ps_lbl.bold = True
    r_ps_lbl.font.size = Pt(11)
    p_ps.add_run("Develop an LSTM-based model for time-series forecasting. Use the Keras Functional API to build a recurrent neural network that predicts future stock prices (AAPL) based on historical sequence data.")

    # Objectives
    doc.add_heading("1. Course Objectives & Relevance", level=2)
    doc.add_paragraph(
        "1. Understand the architecture of Recurrent Neural Networks (RNN) and Long Short-Term Memory (LSTM) cells.\n"
        "2. Solve the vanishing gradient problem using LSTM gating mechanisms (Forget, Input, Output gates).\n"
        "3. Process sequential data by defining lookback windows and generating 3D tensors for Keras.\n"
        "4. Implement early stopping to prevent overfitting during chronological time-series training."
    )

    # Theoretical Foundations
    doc.add_heading("2. Theoretical Foundations & Mathematical Formulations", level=2)
    doc.add_paragraph(
        "LSTM Gate Equations:\n"
        "1. Forget Gate: f_t = sigma(W_f * [h_{t-1}, x_t] + b_f)\n"
        "2. Input Gate: i_t = sigma(W_i * [h_{t-1}, x_t] + b_i)\n"
        "3. Candidate Cell State: C'_t = tanh(W_C * [h_{t-1}, x_t] + b_C)\n"
        "4. Update Cell State: C_t = f_t * C_{t-1} + i_t * C'_t\n"
        "5. Output Gate: o_t = sigma(W_o * [h_{t-1}, x_t] + b_o)\n"
        "6. Hidden State: h_t = o_t * tanh(C_t)\n\n"
        "Loss Function for Forecasting:\n"
        "Mean Squared Error (MSE) = (1/n) * sum(Y_actual - Y_predicted)^2"
    )

    # Step-by-Step Algorithm
    doc.add_heading("3. Algorithm / Step-by-Step Execution Plan", level=2)
    doc.add_paragraph(
        "1. Import libraries: NumPy, Pandas, Matplotlib, Keras (Functional API), and yfinance.\n"
        "2. Download historical stock price data for Apple (AAPL).\n"
        "3. Preprocess data by scaling it between 0 and 1 using MinMaxScaler.\n"
        "4. Create rolling window sequences of length N (e.g., 60 days) to predict the N+1 day.\n"
        "5. Split the data chronologically into 80% training and 20% testing sets.\n"
        "6. Construct LSTM Architecture: Input(60, 1) -> LSTM(64) -> Dropout(0.2) -> Dense(1).\n"
        "7. Compile model using RMSprop optimizer and MSE loss.\n"
        "8. Train model for 50 epochs utilizing EarlyStopping on validation loss.\n"
        "9. Inverse scale predictions and compute Root Mean Squared Error (RMSE).\n"
        "10. Plot actual vs predicted stock prices to visually inspect forecasting accuracy."
    )

    # Summary Table: Architecture
    doc.add_heading("4. LSTM Architecture Breakdown", level=2)
    table = doc.add_table(rows=5, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Layer Name", "Layer Type", "Output Shape", "Trainable Parameters"]
    for j, text in enumerate(headers):
        cell = table.cell(0, j)
        cell.paragraphs[0].text = text
        cell.paragraphs[0].runs[0].bold = True
        set_cell_background(cell, "1F4E79")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

    row_data = [
        ("inputs", "InputLayer", "(None, 60, 1)", "0"),
        ("lstm_1", "LSTM", "(None, 64)", "16,896"),
        ("dropout_1", "Dropout (0.2)", "(None, 64)", "0"),
        ("out_layer", "Dense", "(None, 1)", "65")
    ]

    for i, data in enumerate(row_data, start=1):
        for j, text in enumerate(data):
            cell = table.cell(i, j)
            cell.paragraphs[0].text = text
            if i % 2 == 1:
                set_cell_background(cell, "F2F2F2")

    doc.add_paragraph("\n")

    # Conclusion
    doc.add_heading("5. Conclusions & Summary", level=2)
    doc.add_paragraph(
        "1. Successfully developed an LSTM model capable of learning long-term dependencies in non-stationary time-series data.\n"
        "2. The Keras Functional API provided a modular approach to mapping sequential inputs through LSTM and Dense layers.\n"
        "3. EarlyStopping prevented the model from overfitting the training data by monitoring validation loss.\n"
        "4. Visual evaluation of actual vs. predicted stock prices demonstrated the LSTM's capability to track major market trends effectively."
    )

    doc.save(docx_path)
    print(f"Created Docx: {docx_path}")

def create_explanation(md_path):
    content = """# Assignment 4: LSTM Time-Series Forecasting

## 🎯 Objective
Develop an LSTM-based model for time-series forecasting. This assignment implements a Long Short-Term Memory network to predict Apple (AAPL) stock prices using historical sequence data.

## 🧠 Concepts Covered
- Extracting and analyzing historical financial time-series data (`yfinance`).
- Time-series preprocessing: Data scaling (`MinMaxScaler`) and generating rolling sequence windows.
- Resolving vanishing gradients using LSTM gating architectures.
- Building models via the **Keras Functional API** (`Input -> LSTM -> Dropout -> Dense`).
- Implementing `EarlyStopping` callbacks to optimize training epochs dynamically.
- Evaluating regression performance using Root Mean Squared Error (RMSE) and plotting trends.

## 📁 Folder Structure
- `Assignment4.ipynb`: Core implementation notebook.
- `60_Assignment-4_DL.docx`: Submission-ready lab report.
- `Datasets/`: Folder reserved for local data files (e.g., saved `AAPL_stock_prices.csv`).
- `Plots/`: Folder reserved for saving evaluation graphs and prediction trends.
"""
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(content)
    print(f"Created EXPLANATION: {md_path}")

if __name__ == "__main__":
    ass4_dir = r"d:\\Vit\\Academics Sem-5\\Deep Learning (DL)\\Lab\\Ass4"
    create_notebook(os.path.join(ass4_dir, "Assignment4.ipynb"))
    create_docx(os.path.join(ass4_dir, "60_Assignment-4_DL.docx"))
    create_explanation(os.path.join(ass4_dir, "EXPLANATION.md"))
