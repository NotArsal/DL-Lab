import docx

def create_submission():
    template_path = r'd:\Vit\Academics Sem-5\Deep Learning (DL)\Lab\Ass1\DL_Assignment Submission Format.docx'
    output_path = r'd:\Vit\Academics Sem-5\Deep Learning (DL)\Lab\Ass1\My_Assignment_1_Submission.docx'
    
    doc = docx.Document(template_path)

    algo_text = [
        "1. Start Google Colab and create a new notebook.",
        "2. Import required libraries: TensorFlow, NumPy, and Matplotlib.",
        "3. Verify TensorFlow installation by printing its version.",
        "4. Load the Fashion MNIST dataset using Keras built-in datasets.",
        "5. Print the shapes of the training and testing data to verify correct loading.",
        "6. Define class names for the dataset labels (0-9).",
        "7. Visualize a sample image to understand the pixel intensity values (0-255).",
        "8. Normalize the feature values by dividing image arrays by 255.0 to scale them between 0 and 1.",
        "9. Visualize the first 25 images with their corresponding labels in a grid format.",
        "10. Observe and analyze the output to ensure the data is ready for neural network training."
    ]

    for p in doc.paragraphs:
        if "Problem Statement :" in p.text:
            p.add_run(" Install and configure TensorFlow/Keras in Google Colab. Perform data preprocessing, normalization, train-test split, and visualization on a sample dataset (Fashion MNIST).")
        elif "Algorithm :" in p.text:
            p.add_run("\n" + "\n".join(algo_text))
        elif "Code :" in p.text:
            p.add_run("\n\n[ PLEASE PASTE SCREENSHOTS OF YOUR GOOGLE COLAB CODE HERE ]")
        elif "Output :" in p.text:
            p.add_run("\n\n[ PLEASE PASTE SCREENSHOTS OF YOUR GOOGLE COLAB OUTPUT/GRAPHS HERE ]")

    doc.save(output_path)
    print(f"Successfully generated {output_path}")

if __name__ == "__main__":
    create_submission()
